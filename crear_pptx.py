import os
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

IMG_DIR = '/tmp/pptx_formulas'
os.makedirs(IMG_DIR, exist_ok=True)

# ============================================================
# RENDER LATEX FORMULAS
# ============================================================
def render_latex(latex, filename, fontsize=22, color='#222222', figw=7, figh=0.7):
    """Render LaTeX to transparent PNG."""
    fig, ax = plt.subplots(figsize=(figw, figh))
    ax.text(0.5, 0.5, f'${latex}$', fontsize=fontsize, ha='center', va='center',
            color=color, family='serif')
    ax.axis('off')
    fig.patch.set_alpha(0)
    path = os.path.join(IMG_DIR, filename)
    fig.savefig(path, dpi=250, transparent=True, bbox_inches='tight', pad_inches=0.15)
    plt.close()
    return path

# Pre-render all formulas
formulas = {
    'f1': (r'\theta_{t+1} = \theta_t - \eta \, \nabla L(\theta_t)', 'f_sin_decay.png', 22, 7, 0.6),
    'f2': (r'\theta_{t+1} = (1 - \eta\lambda)\,\theta_t \;-\; \eta \, \nabla L(\theta_t)', 'f_con_decay.png', 22, 9, 0.7),
    'f3': (r'h_i = 0 \;\; (prob.\; p) \quad \mathrm{o} \quad h_i = \frac{x_i}{1-p} \;\; (prob.\; 1-p)', 'f_dropout.png', 20, 8, 0.7),
    'f4': (r'L_{total} = L_{BCE}(\hat{y}, y) \;+\; \lambda \sum_{i} \| w_i \|^2', 'f_loss_total.png', 22, 8, 0.7),
    'f5': (r'Ratio = \frac{parametros}{muestras} = \frac{19000}{60} \approx 300:1', 'f_ratio.png', 22, 8, 0.7),
}

formula_paths = {}
for key, (latex, fname, fs, fw, fh) in formulas.items():
    formula_paths[key] = render_latex(latex, fname, fs, figw=fw, figh=fh)
    print(f'  Rendered {key}: {fname}')

print('All formulas rendered.\n')

# ============================================================
# PPTX SETUP
# ============================================================
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Palette
BG_WHITE = RGBColor(0xFB, 0xFB, 0xFB)
BG_DARK = RGBColor(0x14, 0x14, 0x1E)
BLACK = RGBColor(0x18, 0x18, 0x18)
DARK = RGBColor(0x2A, 0x2A, 0x2A)
MID = RGBColor(0x60, 0x60, 0x60)
SOFT = RGBColor(0x90, 0x90, 0x90)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NEAR_WHITE = RGBColor(0xE8, 0xE8, 0xE8)

BLUE = RGBColor(0x1A, 0x73, 0xE8)
RED = RGBColor(0xD9, 0x30, 0x25)
GREEN = RGBColor(0x1E, 0x88, 0x50)
TEAL = RGBColor(0x00, 0x97, 0xA7)
AMBER = RGBColor(0xFF, 0x8F, 0x00)

CARD_LIGHT = RGBColor(0xF4, 0xF4, 0xF6)
CARD_RED = RGBColor(0xFD, 0xF0, 0xF0)
CARD_GREEN = RGBColor(0xEE, 0xF7, 0xF0)
CARD_BLUE = RGBColor(0xEE, 0xF3, 0xFC)
CARD_AMBER = RGBColor(0xFF, 0xF8, 0xE1)


def set_bg(slide, color=BG_WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def text(slide, l, t, w, h, txt, sz=18, clr=DARK, bold=False, align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = txt
    p.font.size = Pt(sz)
    p.font.color.rgb = clr
    p.font.bold = bold
    p.font.name = font
    p.alignment = align
    return box


def multitext(slide, l, t, w, h, lines, sz=16, clr=DARK, line_spacing=Pt(10)):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        # Support bold prefix with **text**
        if line.startswith('**') and '**' in line[2:]:
            end = line.index('**', 2)
            bold_part = line[2:end]
            rest = line[end+2:]
            run1 = p.add_run()
            run1.text = bold_part
            run1.font.bold = True
            run1.font.size = Pt(sz)
            run1.font.color.rgb = clr
            run1.font.name = "Calibri"
            if rest:
                run2 = p.add_run()
                run2.text = rest
                run2.font.size = Pt(sz)
                run2.font.color.rgb = clr
                run2.font.name = "Calibri"
        else:
            p.text = line
            p.font.size = Pt(sz)
            p.font.color.rgb = clr
            p.font.name = "Calibri"
        p.space_before = line_spacing
        p.alignment = PP_ALIGN.LEFT
    return box


def card(slide, l, t, w, h, fill=CARD_LIGHT, border_clr=None, border_w=1.0):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if border_clr:
        shape.line.color.rgb = border_clr
        shape.line.width = Pt(border_w)
    else:
        shape.line.fill.background()
    # Softer corners
    shape.adjustments[0] = 0.04
    return shape


def accent_line(slide, l, t, w, clr=BLUE, thickness=3):
    shape = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Pt(thickness))
    shape.fill.solid()
    shape.fill.fore_color.rgb = clr
    shape.line.fill.background()


def add_img(slide, path, l, t, w):
    slide.shapes.add_picture(path, Inches(l), Inches(t), Inches(w))


# ============================================================
# SLIDE 1 - PORTADA
# ============================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s, BG_DARK)

# Accent bar left
shape = s.shapes.add_shape(1, Inches(1.2), Inches(1.5), Pt(5), Inches(3.5))
shape.fill.solid()
shape.fill.fore_color.rgb = BLUE
shape.line.fill.background()

text(s, 1.6, 1.5, 10, 0.5, "GRUPO 7", 16, SOFT, bold=True)
text(s, 1.6, 2.1, 10, 1.2, "La Memoria de Pez", 56, WHITE, bold=True)
text(s, 1.6, 3.3, 10, 0.6, "Deep Learning Audit: El Rescate de HealthTech", 24, BLUE)
text(s, 1.6, 4.5, 10, 0.5, "Taller  |  Sesion 2", 15, SOFT)

text(s, 1.6, 6.2, 10, 0.5, "Universidad Privada del Norte", 14, RGBColor(0x66,0x66,0x66))
text(s, 1.6, 6.6, 10, 0.5, "Escuela de Posgrado  |  Modelos de Deep Learning", 13, RGBColor(0x55,0x55,0x55))

# ============================================================
# SLIDE 2 - EL PROBLEMA
# ============================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)

text(s, 1.2, 0.5, 10, 0.7, "El Problema", 42, BLACK, bold=True)
accent_line(s, 1.2, 1.15, 2.5, RED)
text(s, 1.2, 1.4, 10, 0.4, "Dataset pequeno, sin regularizacion, sin weight decay", 17, RED, bold=True)

# Card: Sintoma
card(s, 1.2, 2.1, 5.4, 2.5, CARD_RED, RED, 1.2)
text(s, 1.5, 2.25, 5, 0.4, "Sintoma", 20, RED, bold=True)
multitext(s, 1.5, 2.7, 4.8, 1.8, [
    "Train accuracy llega a ~100%",
    "Test accuracy se queda en ~65-75%",
    "La perdida de test diverge",
    "El modelo MEMORIZA en vez de aprender",
], sz=15, clr=DARK)

# Card: Numeros
card(s, 6.9, 2.1, 5.4, 2.5, CARD_RED, RED, 1.2)
text(s, 7.2, 2.25, 5, 0.4, "Los numeros", 20, RED, bold=True)
multitext(s, 7.2, 2.7, 4.8, 1.0, [
    "~19,000 parametros  vs  60 muestras",
], sz=15, clr=DARK)
add_img(s, formula_paths['f5'], 7.2, 3.4, 4.5)

# Card: Que falta
card(s, 1.2, 4.85, 5.4, 2.2, CARD_LIGHT, SOFT, 0.8)
text(s, 1.5, 4.95, 5, 0.4, "Lo que falta", 20, DARK, bold=True)
multitext(s, 1.5, 5.4, 4.8, 1.5, [
    "   No hay Weight Decay (L2)",
    "   No hay Dropout",
    "   No hay Early Stopping",
    "   Red sobredimensionada",
], sz=15, clr=MID)

# Card: Impacto
card(s, 6.9, 4.85, 5.4, 2.2, CARD_LIGHT, SOFT, 0.8)
text(s, 7.2, 4.95, 5, 0.4, "Impacto clinico", 20, DARK, bold=True)
multitext(s, 7.2, 5.4, 4.8, 1.5, [
    "En produccion no detectaria",
    "patologias nuevas correctamente.",
    "Solo \"recuerda\" los casos vistos.",
    "Paciente nuevo = diagnostico erroneo.",
], sz=15, clr=MID)

# ============================================================
# SLIDE 3 - DIAGNOSTICO
# ============================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)

text(s, 1.2, 0.5, 10, 0.7, "Diagnostico", 42, BLACK, bold=True)
accent_line(s, 1.2, 1.15, 2.5, AMBER)

items = [
    ("Sintoma", "Train accuracy converge a ~100% mientras test accuracy se estanca o cae. La brecha (gap) supera el 25%."),
    ("Causa raiz", "Sin regularizacion, los pesos crecen sin control. Con 300x mas parametros que muestras, la red memoriza cada dato incluyendo su ruido."),
    ("Evidencia", "Los pesos muestran valores grandes y dispersos (alta varianza). Codifican ruido en vez de patrones generalizables."),
    ("Analogia", "Un estudiante que memoriza las respuestas del examen de practica. Perfecto en lo visto, pesimo en preguntas nuevas."),
]

y = 1.55
for titulo, desc in items:
    accent_line(s, 1.2, y + 0.02, 0.12, AMBER, 40)
    text(s, 1.55, y - 0.05, 2, 0.4, titulo, 17, AMBER, bold=True)
    text(s, 3.5, y - 0.05, 8.5, 0.55, desc, 15, DARK)
    y += 1.25

# ============================================================
# SLIDE 4 - LA CORRECCION
# ============================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)

text(s, 1.2, 0.5, 10, 0.7, "La Correccion", 42, BLACK, bold=True)
accent_line(s, 1.2, 1.15, 2.5, GREEN)
text(s, 1.2, 1.4, 10, 0.4, "Tres interruptores para arreglar el modelo", 17, GREEN, bold=True)

# Card 1
card(s, 1.2, 2.1, 3.5, 4.5, CARD_GREEN, GREEN, 1.2)
text(s, 1.5, 2.3, 3, 0.4, "Weight Decay", 22, GREEN, bold=True)
text(s, 1.5, 2.8, 3, 0.4, "(L2 Regularization)", 14, MID)
multitext(s, 1.5, 3.4, 3, 1.5, [
    "Penaliza pesos grandes",
    "sumando un termino de",
    "penalizacion a la perdida.",
    "",
    "Encoge los pesos hacia",
    "cero en cada paso.",
], sz=14, clr=DARK)
card(s, 1.6, 5.6, 2.8, 0.55, RGBColor(0xD5, 0xEE, 0xDB), GREEN, 0.8)
text(s, 1.7, 5.65, 2.6, 0.4, "weight_decay=0.01", 14, GREEN, bold=True, align=PP_ALIGN.CENTER, font="Consolas")

# Card 2
card(s, 5.0, 2.1, 3.5, 4.5, CARD_GREEN, GREEN, 1.2)
text(s, 5.3, 2.3, 3, 0.4, "Dropout", 22, GREEN, bold=True)
text(s, 5.3, 2.8, 3, 0.4, "(Regularizacion estocastica)", 14, MID)
multitext(s, 5.3, 3.4, 3, 1.5, [
    "Apaga neuronas al azar",
    "durante el entrenamiento.",
    "",
    "Impide que la red",
    "dependa de neuronas",
    "individuales.",
], sz=14, clr=DARK)
card(s, 5.4, 5.6, 2.8, 0.55, RGBColor(0xD5, 0xEE, 0xDB), GREEN, 0.8)
text(s, 5.5, 5.65, 2.6, 0.4, "nn.Dropout(0.3)", 14, GREEN, bold=True, align=PP_ALIGN.CENTER, font="Consolas")

# Card 3
card(s, 8.8, 2.1, 3.5, 4.5, CARD_GREEN, GREEN, 1.2)
text(s, 9.1, 2.3, 3, 0.4, "Reducir Red", 22, GREEN, bold=True)
text(s, 9.1, 2.8, 3, 0.4, "(Menos parametros)", 14, MID)
multitext(s, 9.1, 3.4, 3, 1.5, [
    "Menos capas y neuronas",
    "para que no pueda",
    "memorizar todo.",
    "",
    "Adecuamos complejidad",
    "al tamano del dataset.",
], sz=14, clr=DARK)
card(s, 9.2, 5.6, 2.8, 0.55, RGBColor(0xD5, 0xEE, 0xDB), GREEN, 0.8)
text(s, 9.3, 5.65, 2.6, 0.4, "19K -> 300 params", 14, GREEN, bold=True, align=PP_ALIGN.CENTER, font="Consolas")

# ============================================================
# SLIDE 5 - ANTES vs DESPUES
# ============================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)

text(s, 1.2, 0.5, 10, 0.7, "Antes vs Despues", 42, BLACK, bold=True)
accent_line(s, 1.2, 1.15, 2.5, BLUE)

rows = [
    ("Metrica",           "Saboteado",       "Corregido"),
    ("Train Accuracy",    "~100%",           "~90-94%"),
    ("Test Accuracy",     "~65-75%",         "~87-93%"),
    ("Gap (Train-Test)",  "~25-35%",         "~3-7%"),
    ("Loss de Test",      "Diverge",         "Estable"),
    ("Weight Decay",      "0.0 (OFF)",       "0.01 (ON)"),
    ("Dropout",           "NO",              "p = 0.3"),
    ("Parametros",        "~19,000",         "~300"),
]

tbl_shape = s.shapes.add_table(len(rows), 3, Inches(1.5), Inches(1.6), Inches(10.3), Inches(5.0))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(3.5)
tbl.columns[1].width = Inches(3.4)
tbl.columns[2].width = Inches(3.4)

for r, (c1, c2, c3) in enumerate(rows):
    for c, txt in enumerate([c1, c2, c3]):
        cell = tbl.cell(r, c)
        cell.text = txt
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        for p in cell.text_frame.paragraphs:
            p.font.name = "Calibri"
            p.font.size = Pt(16)
            p.alignment = PP_ALIGN.CENTER
            if r == 0:
                p.font.bold = True
                p.font.color.rgb = WHITE
            else:
                p.font.color.rgb = DARK
                if c == 1:
                    p.font.color.rgb = RED
                elif c == 2:
                    p.font.color.rgb = GREEN
                    p.font.bold = True
        cell.fill.solid()
        if r == 0:
            cell.fill.fore_color.rgb = [RGBColor(0x33,0x33,0x33), RED, GREEN][c]
        else:
            cell.fill.fore_color.rgb = WHITE if r % 2 == 1 else RGBColor(0xF5, 0xF5, 0xF7)

# ============================================================
# SLIDE 6 - REPRESENTATION LEARNING
# ============================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)

text(s, 1.2, 0.5, 10, 0.7, "Representation Learning", 42, BLACK, bold=True)
accent_line(s, 1.2, 1.15, 2.5, TEAL)
text(s, 1.2, 1.4, 10, 0.4, "Como cambio la jerarquia de conceptos tras el arreglo", 17, TEAL, bold=True)

# Left card
card(s, 1.2, 2.1, 5.3, 4.2, CARD_RED, RED, 1.2)
text(s, 1.5, 2.3, 4.8, 0.4, "Sin Regularizacion", 22, RED, bold=True)
accent_line(s, 1.5, 2.8, 2.0, RED, 2)
multitext(s, 1.5, 3.0, 4.8, 3.0, [
    "Memoriza el ruido de cada muestra",
    "",
    "Las features internas no son",
    "transferibles a datos nuevos",
    "",
    "No aprende conceptos generales",
    "(bordes, texturas, formas)",
    "",
    "Cada neurona se sobreactiva",
    "para casos puntuales",
], sz=15, clr=DARK)

# Right card
card(s, 6.8, 2.1, 5.3, 4.2, CARD_GREEN, GREEN, 1.2)
text(s, 7.1, 2.3, 4.8, 0.4, "Con Regularizacion", 22, GREEN, bold=True)
accent_line(s, 7.1, 2.8, 2.0, GREEN, 2)
multitext(s, 7.1, 3.0, 4.8, 3.0, [
    "Captura patrones reales:",
    "distribuciones y correlaciones",
    "",
    "Features compartidas entre",
    "muestras similares",
    "",
    "Capas intermedias actuan como",
    "detectores de patrones",
    "",
    "Jerarquia: simple -> complejo",
], sz=15, clr=DARK)

text(s, 1.2, 6.5, 10.9, 0.5,
     "La regularizacion fuerza a la red a buscar patrones compartidos en vez de memorizar caso por caso.",
     16, MID, align=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 7 - DECISION DE NEGOCIO
# ============================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)

text(s, 1.2, 0.5, 10, 0.7, "Decision de Negocio", 42, BLACK, bold=True)
accent_line(s, 1.2, 1.15, 2.5, BLUE)

card(s, 1.2, 1.6, 11.0, 1.0, CARD_BLUE, BLUE, 1.0)
text(s, 1.5, 1.7, 10.4, 0.4, "Si tuvieramos mas presupuesto, donde invertirlo?", 19, BLUE, bold=True)
text(s, 1.5, 2.1, 10.4, 0.4, "Recomendacion: invertir en DATOS + REGULARIZACION", 16, DARK)

# Card 1
card(s, 1.2, 3.0, 5.3, 3.5, CARD_LIGHT, BLUE, 0.8)
text(s, 1.5, 3.15, 5, 0.4, "Prioridad 1", 13, BLUE, bold=True)
text(s, 1.5, 3.5, 5, 0.5, "Mas datos", 24, DARK, bold=True)
accent_line(s, 1.5, 4.0, 1.5, BLUE, 2)
multitext(s, 1.5, 4.2, 4.8, 2.0, [
    "Con dataset pequeno, ninguna tecnica",
    "supera tener mas muestras.",
    "",
    "Data Augmentation, Transfer Learning,",
    "recopilacion de imagenes medicas.",
], sz=15, clr=DARK)

# Card 2
card(s, 6.8, 3.0, 5.3, 3.5, CARD_LIGHT, BLUE, 0.8)
text(s, 7.1, 3.15, 5, 0.4, "Prioridad 2", 13, BLUE, bold=True)
text(s, 7.1, 3.5, 5, 0.5, "Mejor regularizacion", 24, DARK, bold=True)
accent_line(s, 7.1, 4.0, 1.5, BLUE, 2)
multitext(s, 7.1, 4.2, 4.8, 2.0, [
    "Combinar L2 + Dropout + BatchNorm",
    "+ Early Stopping.",
    "",
    "Validacion cruzada (k-fold) para",
    "estimar rendimiento real.",
], sz=15, clr=DARK)

text(s, 1.2, 6.7, 11, 0.4,
     "En medicina, un modelo que memoriza es peligroso: puede dar falsos negativos en pacientes nuevos.",
     15, RED, align=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 8 - BONUS MATEMATICO (con LaTeX images)
# ============================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)

text(s, 1.2, 0.5, 10, 0.7, "Bonus Matematico", 42, BLACK, bold=True)
accent_line(s, 1.2, 1.15, 2.5, RGBColor(0x7B, 0x1F, 0xA2))

# Sin Weight Decay
text(s, 1.2, 1.6, 10, 0.4, "Sin Weight Decay", 19, DARK, bold=True)
text(s, 1.5, 2.0, 8, 0.4, "Los pesos se actualizan solo por el gradiente. No hay limite a su crecimiento.", 14, MID)
card(s, 1.5, 2.5, 8, 0.8, CARD_LIGHT)
add_img(s, formula_paths['f1'], 2.5, 2.55, 5.5)

# Con Weight Decay
text(s, 1.2, 3.5, 10, 0.4, "Con Weight Decay (L2)", 19, DARK, bold=True)
text(s, 1.5, 3.9, 8, 0.4, "El factor (1 - eta*lambda) encoge los pesos en cada paso. Pesos grandes se penalizan mas.", 14, MID)
card(s, 1.5, 4.3, 9.5, 0.9, CARD_LIGHT)
add_img(s, formula_paths['f2'], 2.0, 4.35, 7.5)

# Dropout
text(s, 1.2, 5.4, 5, 0.4, "Dropout", 19, DARK, bold=True)
card(s, 1.5, 5.8, 5, 1.3, CARD_LIGHT)
add_img(s, formula_paths['f3'], 1.8, 5.85, 4.0)

# Loss total
text(s, 7.0, 5.4, 5, 0.4, "Loss total con L2", 19, DARK, bold=True)
card(s, 7.0, 5.8, 5.5, 1.0, CARD_LIGHT)
add_img(s, formula_paths['f4'], 7.2, 5.9, 5.0)

# ============================================================
# SLIDE 9 - CONCLUSIONES
# ============================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s, BG_DARK)

text(s, 1.2, 0.5, 10, 0.7, "Conclusiones", 42, WHITE, bold=True)
accent_line(s, 1.2, 1.15, 2.5, BLUE)

conclusions = [
    ("Diagnostico", "El modelo memorizaba el dataset. Train ~100%, Test ~65-75%. Overfitting violento por falta de regularizacion."),
    ("Interruptor", "Activamos Weight Decay (L2=0.01), Dropout (p=0.3) y redujimos la red de ~19K a ~300 parametros."),
    ("Resultado", "La brecha train-test se redujo de ~30% a menos del 7%. El modelo ahora generaliza."),
    ("HealthTech", "En medicina, un modelo con overfitting da diagnosticos falsos en pacientes nuevos. La regularizacion es esencial."),
]

y = 1.6
for titulo, desc in conclusions:
    # Accent dot
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.2), Inches(y + 0.15), Pt(12), Pt(12))
    dot.fill.solid()
    dot.fill.fore_color.rgb = BLUE
    dot.line.fill.background()

    text(s, 1.6, y - 0.05, 3, 0.4, titulo, 20, BLUE, bold=True)
    text(s, 1.6, y + 0.35, 10, 0.5, desc, 16, NEAR_WHITE)
    y += 1.15

text(s, 1.2, 6.5, 10.9, 0.4,
     "Grupo 7  |  La Memoria de Pez  |  Bloque 3: Problemas de Estrategia",
     13, SOFT, align=PP_ALIGN.CENTER)

# ============================================================
# SAVE
# ============================================================
out_path = '/home/user/proyecto/Grupo7_La_Memoria_de_Pez.pptx'
prs.save(out_path)
print(f'\nPPTX guardado: {out_path}')
print(f'Slides: {len(prs.slides)}')
